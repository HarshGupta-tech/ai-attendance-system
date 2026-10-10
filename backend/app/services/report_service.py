import csv
import io
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import openpyxl
from sqlalchemy.orm import Session
from typing import List, Dict, Optional
from app.models import Student, AttendanceRecord, AttendanceSession, Subject

class ReportService:
    def __init__(self, db: Session):
        self.db = db
    
    def generate_attendance_report(self, filters: dict) -> List[dict]:
        """Generate attendance report data based on filters."""
        query = self.db.query(AttendanceRecord, Student, AttendanceSession, Subject)\
            .join(Student)\
            .join(AttendanceSession)\
            .join(Subject)
            
        if filters.get('start_date'):
            query = query.filter(AttendanceSession.date >= filters['start_date'])
        if filters.get('end_date'):
            query = query.filter(AttendanceSession.date <= filters['end_date'])
        if filters.get('subject_id'):
            query = query.filter(AttendanceSession.subject_id == filters['subject_id'])
            
        results = query.all()
        report = []
        for record, student, session, subject in results:
            report.append({
                "student_id": student.student_id,
                "student_name": student.full_name,
                "subject_code": subject.code,
                "date": session.date.isoformat(),
                "status": record.status
            })
        return report
    
    def generate_student_report(self, student_id: Optional[int] = None,
                                 course_id: Optional[int] = None,
                                 semester: Optional[int] = None) -> List[dict]:
        """Generate student-wise attendance report."""
        query = self.db.query(Student)
        if student_id:
            query = query.filter(Student.id == student_id)
        if course_id:
            query = query.filter(Student.course_id == course_id)
        if semester:
            query = query.filter(Student.semester == semester)
            
        students = query.all()
        report = []
        for student in students:
            total = self.db.query(AttendanceRecord).filter(AttendanceRecord.student_id == student.id).count()
            present = self.db.query(AttendanceRecord).filter(
                AttendanceRecord.student_id == student.id,
                AttendanceRecord.status.in_(['present', 'late'])
            ).count()
            percentage = (present / total * 100) if total > 0 else 0
            
            report.append({
                "student_id": student.student_id,
                "student_name": student.full_name,
                "total_sessions": total,
                "attended": present,
                "percentage": round(percentage, 2)
            })
        return report
    
    def generate_subject_report(self, subject_id: Optional[int] = None) -> List[dict]:
        query = self.db.query(Subject)
        if subject_id:
            query = query.filter(Subject.id == subject_id)
            
        subjects = query.all()
        report = []
        for subject in subjects:
            total = self.db.query(AttendanceRecord).join(AttendanceSession)\
                .filter(AttendanceSession.subject_id == subject.id).count()
            present = self.db.query(AttendanceRecord).join(AttendanceSession)\
                .filter(AttendanceSession.subject_id == subject.id, AttendanceRecord.status.in_(['present', 'late'])).count()
            percentage = (present / total * 100) if total > 0 else 0
            
            report.append({
                "subject_code": subject.code,
                "subject_name": subject.name,
                "total_records": total,
                "present_records": present,
                "percentage": round(percentage, 2)
            })
        return report
    
    def generate_low_attendance_report(self, threshold: float = 75.0) -> List[dict]:
        students = self.db.query(Student).all()
        report = []
        for student in students:
            total = self.db.query(AttendanceRecord).filter(AttendanceRecord.student_id == student.id).count()
            if total == 0: continue
            present = self.db.query(AttendanceRecord).filter(
                AttendanceRecord.student_id == student.id,
                AttendanceRecord.status.in_(['present', 'late'])
            ).count()
            percentage = (present / total * 100)
            if percentage < threshold:
                report.append({
                    "student_id": student.student_id,
                    "student_name": student.full_name,
                    "total_sessions": total,
                    "attended": present,
                    "percentage": round(percentage, 2)
                })
        return report
    
    def export_csv(self, data: List[dict], columns: List[str]) -> str:
        """Export data to CSV string."""
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=columns)
        writer.writeheader()
        for row in data:
            writer.writerow({k: row.get(k, '') for k in columns})
        return output.getvalue()
    
    def export_excel(self, data: List[dict], columns: List[str], title: str = 'Report') -> bytes:
        """Export data to Excel bytes."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title
        
        # Header styling
        from openpyxl.styles import Font, PatternFill, Alignment
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        
        # Write headers
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name.replace('_', ' ').title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
        
        # Write data
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, col_name in enumerate(columns, 1):
                ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ''))
        
        # Auto-fit column widths
        for col_idx, col_name in enumerate(columns, 1):
            max_len = max(len(str(col_name)), max((len(str(row.get(col_name, ''))) for row in data), default=0))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 30)
        
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
    
    def export_pdf(self, data: List[dict], columns: List[str], title: str = 'Report') -> bytes:
        """Export data to PDF bytes."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
        
        # Title
        elements.append(Paragraph(title, styles['Title']))
        elements.append(Spacer(1, 20))
        
        # Subtitle with generation info
        from datetime import datetime
        elements.append(Paragraph(
            f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            styles['Normal']
        ))
        elements.append(Spacer(1, 20))
        
        # Table
        headers = [col.replace('_', ' ').title() for col in columns]
        table_data = [headers]
        for row in data:
            table_data.append([str(row.get(col, '')) for col in columns])
        
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        
        elements.append(table)
        doc.build(elements)
        return output.getvalue()
